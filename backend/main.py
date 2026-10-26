from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
import uvicorn
import os
import shutil
from pathlib import Path
import fitz  # PyMuPDF
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import io
import csv
from PIL import Image
import pytesseract
import re
import arabic_reshaper
from bidi.algorithm import get_display
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Configure Tesseract OCR
TESSERACT_EXE = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
if TESSERACT_EXE.exists():
    pytesseract.pytesseract.tesseract_cmd = str(TESSERACT_EXE)

# Register Unicode fonts for Arabic and Multilingual support in ReportLab
ARIAL_FONT_PATH = Path("C:/Windows/Fonts/arial.ttf")
ARIAL_BOLD_PATH = Path("C:/Windows/Fonts/arialbd.ttf")

if ARIAL_FONT_PATH.exists():
    pdfmetrics.registerFont(TTFont('ArabicFont', str(ARIAL_FONT_PATH)))
    pdfmetrics.registerFont(TTFont('ArabicFont-Bold', str(ARIAL_BOLD_PATH if ARIAL_BOLD_PATH.exists() else ARIAL_FONT_PATH)))
    DEFAULT_FONT = 'ArabicFont'
    DEFAULT_BOLD_FONT = 'ArabicFont-Bold'
else:
    DEFAULT_FONT = 'Helvetica'
    DEFAULT_BOLD_FONT = 'Helvetica-Bold'


def prepare_arabic_text(text: str) -> str:
    """Reshape Arabic text and apply bidirectional algorithm so it displays correctly in PDF."""
    if not text:
        return ""
    # Check if string contains Arabic characters
    if re.search(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]', text):
        try:
            reshaped = arabic_reshaper.reshape(text)
            return get_display(reshaped)
        except Exception:
            return text
    return text

app = FastAPI(title="SmartConvert AI", description="Intelligent File Conversion API")

# Enable CORS for frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Directory for temporary files
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


def ocr_page(page) -> str:
    """Perform OCR on a PDF page image if text is not selectable."""
    try:
        pix = page.get_pixmap(dpi=150)
        img = Image.open(io.BytesIO(pix.tobytes()))
        text = pytesseract.image_to_string(img, lang="ara+eng")
        return text.strip()
    except Exception as e:
        print(f"OCR warning: {e}")
        return ""


def get_page_text_with_ocr(page) -> str:
    """Extract text from a PDF page, falling back to OCR for scanned pages."""
    text = page.get_text().strip()
    if not text:
        text = ocr_page(page)
    return text


@app.post("/uploadfile/")
async def create_upload_file(file: UploadFile = File(...)):
    """Upload a file and return basic information."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file uploaded")
    
    file_path = UPLOAD_DIR / file.filename
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    finally:
        file.file.close()
    
    return JSONResponse(
        content={
            "filename": file.filename,
            "content_type": file.content_type,
            "size": file_path.stat().st_size,
            "message": "File uploaded successfully"
        }
    )


@app.post("/convert/")
async def convert_file(
    file: UploadFile = File(...),
    target_format: str = Form(...)
):
    """Convert an uploaded file to the target format."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file uploaded")
    
    file_path = UPLOAD_DIR / file.filename
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    finally:
        file.file.close()
    
    file_ext = file_path.suffix.lower()
    target_format = target_format.lower()
    
    # PDF conversions
    if file_ext == ".pdf":
        if target_format == "txt":
            return await pdf_to_txt(file_path)
        elif target_format == "csv":
            return await pdf_to_csv(file_path)
        elif target_format == "docx":
            return await pdf_to_docx(file_path)
        elif target_format == "xlsx":
            return await pdf_to_xlsx(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for PDF")
    
    # Text conversions
    elif file_ext == ".txt":
        if target_format == "pdf":
            return await txt_to_pdf(file_path)
        elif target_format == "docx":
            return await txt_to_docx(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for TXT")
    
    # Image conversions
    elif file_ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
        if target_format == "pdf":
            return await image_to_pdf(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for image")
    
    # Excel conversions
    elif file_ext in [".xlsx", ".xls"]:
        if target_format == "pdf":
            return await excel_to_pdf(file_path)
        elif target_format == "docx":
            return await excel_to_docx(file_path)
        elif target_format == "csv":
            return await excel_to_csv(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for Excel")
    
    # Word conversions
    elif file_ext == ".docx":
        if target_format == "pdf":
            return await docx_to_pdf(file_path)
        elif target_format == "xlsx":
            return await docx_to_xlsx(file_path)
        elif target_format == "txt":
            return await docx_to_txt(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for Word")
    
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format {file_ext}")


async def pdf_to_txt(file_path: Path):
    """Extract text from PDF (with OCR fallback) and return as text file."""
    doc = fitz.open(file_path)
    pages_text = []
    for page in doc:
        pages_text.append(get_page_text_with_ocr(page))
    doc.close()
    
    full_text = "\n\n--- Page Break ---\n\n".join(pages_text)
    
    output = io.BytesIO()
    output.write(full_text.encode('utf-8'))
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.txt"}
    )


async def pdf_to_docx(file_path: Path):
    """Convert PDF to Word document with OCR fallback."""
    doc = Document()
    pdf_doc = fitz.open(file_path)
    
    for i, page in enumerate(pdf_doc):
        if i > 0:
            doc.add_page_break()
        
        # Check for native tables
        tables = list(page.find_tables())
        if tables:
            for t in tables:
                rows = t.extract()
                if rows:
                    word_table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    word_table.style = 'Table Grid'
                    for r_idx, row in enumerate(rows):
                        for c_idx, val in enumerate(row):
                            word_table.cell(r_idx, c_idx).text = str(val or "")
        
        text = get_page_text_with_ocr(page)
        for line in text.split('\n'):
            line_str = line.strip()
            if line_str:
                doc.add_paragraph(line_str)
                
    pdf_doc.close()
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.docx"}
    )


async def pdf_to_xlsx(file_path: Path):
    """Extract tables and text from PDF (with OCR fallback) into a single,
    well-formatted Excel sheet.  Every page gets a coloured header banner,
    its content in clean rows, then a blank spacer before the next page.
    """
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter

    # ── colour palette ─────────────────────────────────────────────────────
    HEADER_FILL   = PatternFill("solid", fgColor="1F3864")   # dark navy
    HEADER_FONT   = Font(bold=True, color="FFFFFF", size=13)
    ROW_ODD_FILL  = PatternFill("solid", fgColor="DCE6F1")   # light blue
    ROW_EVN_FILL  = PatternFill("solid", fgColor="FFFFFF")   # white
    ROW_FONT      = Font(size=11)
    THIN_SIDE     = Side(style="thin", color="B8CCE4")
    THIN_BORDER   = Border(left=THIN_SIDE, right=THIN_SIDE,
                           top=THIN_SIDE,  bottom=THIN_SIDE)
    WRAP_ALIGN    = Alignment(wrap_text=True, vertical="center",
                              horizontal="right")   # RTL-friendly
    CENTER_ALIGN  = Alignment(horizontal="center", vertical="center")

    pdf_doc = fitz.open(file_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Content"
    ws.sheet_view.rightToLeft = True      # RTL sheet for Arabic

    # Freeze the top row so headers stay visible while scrolling
    ws.freeze_panes = "A2"

    current_row = 1

    for page_idx, page in enumerate(pdf_doc):
        page_num = page_idx + 1

        # ── Page header banner ──────────────────────────────────────────────
        ws.row_dimensions[current_row].height = 24
        cell = ws.cell(row=current_row, column=1,
                       value=f"  صفحة {page_num}  •  Page {page_num}")
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        # Merge across 3 columns for the banner
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row,   end_column=3)
        current_row += 1

        # ── Content ─────────────────────────────────────────────────────────
        tables_found = list(page.find_tables())
        content_rows: list[list[str]] = []

        if tables_found:
            for t in tables_found:
                for trow in t.extract():
                    content_rows.append([str(c or "").strip() for c in trow])
                content_rows.append([])          # blank between tables
        else:
            text = get_page_text_with_ocr(page)
            for line in text.split("\n"):
                clean = line.strip()
                if clean:
                    content_rows.append([clean])

        # Write rows with alternating fill + borders
        for rel_idx, data_row in enumerate(content_rows):
            if not any(c for c in data_row):
                current_row += 1
                continue

            fill = ROW_ODD_FILL if rel_idx % 2 == 0 else ROW_EVN_FILL
            ws.row_dimensions[current_row].height = 18

            for col_idx, value in enumerate(data_row, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=value)
                c.fill      = fill
                c.font      = ROW_FONT
                c.border    = THIN_BORDER
                c.alignment = WRAP_ALIGN

            current_row += 1

        # ── Spacer row between pages ────────────────────────────────────────
        ws.row_dimensions[current_row].height = 8
        current_row += 2

    pdf_doc.close()

    # ── Auto-fit column widths (max 80 chars) ────────────────────────────────
    col_widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and not isinstance(cell, type(None)):
                col_widths[cell.column] = max(
                    col_widths.get(cell.column, 0),
                    min(len(str(cell.value)), 80)
                )
    for col_num, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width + 4

    # Ensure at least 3 columns have a reasonable minimum width
    for col_num in range(1, 4):
        ltr = get_column_letter(col_num)
        if ws.column_dimensions[ltr].width < 20:
            ws.column_dimensions[ltr].width = 20

    # ── Output ───────────────────────────────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.xlsx"}
    )



async def pdf_to_csv(file_path: Path):
    """Extract tables or text from PDF and return as CSV."""
    pdf_doc = fitz.open(file_path)
    all_rows = []
    
    for page in pdf_doc:
        tables = list(page.find_tables())
        if tables:
            for t in tables:
                for row in t.extract():
                    all_rows.append([str(c or "") for c in row])
        else:
            text = get_page_text_with_ocr(page)
            for line in text.split('\n'):
                line_str = line.strip()
                if line_str:
                    all_rows.append([line_str])
    
    pdf_doc.close()
    
    output = io.StringIO()
    writer = csv.writer(output)
    for row in all_rows:
        writer.writerow(row)
    
    bytes_output = io.BytesIO(output.getvalue().encode('utf-8-sig'))
    bytes_output.seek(0)
    
    return StreamingResponse(
        bytes_output,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.csv"}
    )


async def txt_to_pdf(file_path: Path):
    """Convert text file to PDF with Arabic support."""
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()
    
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    style = styles["Normal"]
    style.fontName = DEFAULT_FONT
    style.fontSize = 11
    style.leading = 16
    
    paragraphs = text.split('\n')
    flowables = []
    for para in paragraphs:
        if para.strip():
            flowables.append(Paragraph(prepare_arabic_text(para.strip()), style))
    
    if not flowables:
        flowables = [Paragraph("Empty file", style)]
        
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )


async def txt_to_docx(file_path: Path):
    """Convert text file to Word document."""
    doc = Document()
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            line_str = line.strip()
            if line_str:
                doc.add_paragraph(line_str)
                
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.docx"}
    )


async def image_to_pdf(file_path: Path):
    """Convert image to PDF."""
    img = Image.open(file_path)
    if img.mode != 'RGB':
        img = img.convert('RGB')
    
    output = io.BytesIO()
    img.save(output, format='PDF')
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )


async def excel_to_pdf(file_path: Path):
    """Convert Excel to clean, centered document paragraphs in PDF (not raw grid table)."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    
    output = io.BytesIO()
    pdf = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=48,
        rightMargin=48,
        topMargin=40,
        bottomMargin=40
    )
    styles = getSampleStyleSheet()
    
    # Text styles
    body_style = styles["Normal"].clone('DocBodyStyle')
    body_style.fontName = DEFAULT_FONT
    body_style.fontSize = 11
    body_style.leading = 18
    body_style.alignment = 1  # Centered
    body_style.textColor = colors.HexColor('#1A202C')
    
    banner_text_style = styles["Normal"].clone('BannerTextStyle')
    banner_text_style.fontName = DEFAULT_BOLD_FONT
    banner_text_style.fontSize = 12
    banner_text_style.leading = 16
    banner_text_style.alignment = 1  # Centered
    banner_text_style.textColor = colors.whitesmoke
    
    flowables = []
    first_section = True
    
    for row in ws.iter_rows(values_only=True):
        non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if not non_empty:
            continue
            
        full_line = "  •  ".join(non_empty) if len(non_empty) > 1 else non_empty[0]
        
        # Check if line is a page header banner
        is_page_header = "صفحة" in full_line or "Page" in full_line
        
        if is_page_header:
            if not first_section:
                flowables.append(Spacer(1, 14))
            first_section = False
            
            # Create a nice centered banner pill
            reshaped_banner = prepare_arabic_text(full_line)
            banner_p = Paragraph(f"<b>{reshaped_banner}</b>", banner_text_style)
            
            # 516 pt available width (612 - 96)
            banner_table = Table([[banner_p]], colWidths=[516])
            banner_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#1F3864')),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('LEFTPADDING', (0, 0), (-1, -1), 12),
                ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ]))
            flowables.append(banner_table)
            flowables.append(Spacer(1, 8))
        else:
            reshaped_text = prepare_arabic_text(full_line)
            p = Paragraph(reshaped_text, body_style)
            flowables.append(p)
            flowables.append(Spacer(1, 5))
            
    if not flowables:
        flowables = [Paragraph(prepare_arabic_text("لا توجد بيانات داخل ملف الإكسيل"), body_style)]
        
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )


async def excel_to_docx(file_path: Path):
    """Convert Excel sheet to Word document table."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    doc = Document()
    
    rows_data = []
    for row in ws.iter_rows(values_only=True):
        rows_data.append([str(c) if c is not None else '' for c in row])
        
    if rows_data:
        max_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=max_cols)
        table.style = 'Table Grid'
        for r_idx, row in enumerate(rows_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    else:
        doc.add_paragraph("Empty Excel spreadsheet.")
        
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.docx"}
    )


async def excel_to_csv(file_path: Path):
    """Convert Excel sheet to CSV."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    
    output = io.StringIO()
    writer = csv.writer(output)
    for row in ws.iter_rows(values_only=True):
        writer.writerow([str(c) if c is not None else '' for c in row])
        
    bytes_output = io.BytesIO(output.getvalue().encode('utf-8-sig'))
    bytes_output.seek(0)
    
    return StreamingResponse(
        bytes_output,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.csv"}
    )


async def docx_to_pdf(file_path: Path):
    """Convert Word to PDF with Arabic support."""
    doc = Document(file_path)
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    style = styles["Normal"]
    style.fontName = DEFAULT_FONT
    style.fontSize = 11
    style.leading = 16
    
    full_text = [p.text for p in doc.paragraphs if p.text.strip()]
    flowables = [Paragraph(prepare_arabic_text(p), style) for p in full_text]
    
    if not flowables:
        flowables = [Paragraph("Empty document", style)]
        
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )


async def docx_to_xlsx(file_path: Path):
    """Convert Word document paragraphs and tables to Excel."""
    doc = Document(file_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Document Content"
    
    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            ws.append([cell.text.strip() for cell in row.cells])
        ws.append([])  # Spacer row
        
    # Extract paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            ws.append([p.text.strip()])
            
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.xlsx"}
    )


async def docx_to_txt(file_path: Path):
    """Convert Word document to plain text."""
    doc = Document(file_path)
    lines = [p.text for p in doc.paragraphs]
    
    output = io.BytesIO("\n".join(lines).encode('utf-8'))
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.txt"}
    )


@app.get("/")
async def root():
    return {"message": "SmartConvert AI API is running"}


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)